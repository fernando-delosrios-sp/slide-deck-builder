#!/usr/bin/env python3
"""
Modernization coverage gate.

Verifies that a modernized spec.yaml faithfully accounts for every slide and
every meaningful visual of its baseline .pptx, and surfaces content that does
not come from the baseline.

Usage:
    check_coverage.py <spec.yaml> [baseline.pptx] [--strict]

The baseline path defaults to the spec's top-level `baseline` key (relative to
the spec file). Exit code 1 when any ERROR is found (also on warnings with
--strict). Run it before every build; deck-builder refuses to build on failure.

Checks
  ERROR  baseline slide neither carried (slide.baseline_slides) nor cut
         (baseline_cuts with slide + reason)
  ERROR  spec slide with no baseline_slides and no `new: true`
  ERROR  significant baseline media (raster >=150px both sides) that is not
         embedded in the spec, not listed in a slide's `redrawn_media`, and
         not in baseline_cuts
  ERROR  baseline_cuts entry without a reason
  WARN   diagram slide carried without an embedded image, diagram or redraw
  WARN   capitalized terms in the spec that never occur in the baseline text
         and are not declared in that slide's `additions`
  INFO   every declared `additions` entry (to show the user at the gate)
"""

import os
import re
import sys
import tempfile

import xml.etree.ElementTree as ET
import yaml
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import inspect_pptx  # noqa: E402

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

MIN_SIDE = 150
STOP = {
    "The", "This", "That", "These", "Those", "With", "From", "When", "Each", "Every",
    "Once", "Only", "Both", "Never", "Not", "And", "For", "Can", "Use", "Define",
    "Identity", "Account", "Access", "Source", "Role", "Roles",
}


def walk_text(shape, out):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            walk_text(s, out)
    elif getattr(shape, "has_text_frame", False):
        out.append(shape.text_frame.text)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                out.append(cell.text)


def is_slide_hidden(slide):
    """Returns True if the slide is marked as hidden in PowerPoint (show='0' or show='false')."""
    show = slide._element.get("show")
    return show is not None and show in ("0", "false")


def baseline_model(pptx_path, include_hidden=False):
    prs = pptx.Presentation(pptx_path)
    corpus, slides = [], {}
    hidden_slides = set()
    canvas_media = set()
    with tempfile.TemporaryDirectory() as tmp:
        media_map, files = inspect_pptx.extract_media_and_rels(pptx_path, tmp)
        sizes = {}
        for name, path in files.items():
            if Image and name.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp")):
                try:
                    with Image.open(path) as im:
                        sizes[name] = im.size
                except Exception:
                    pass

        sw = prs.slide_width
        sh = prs.slide_height
        for idx, slide in enumerate(prs.slides, 1):
            is_hidden = is_slide_hidden(slide)
            if is_hidden:
                hidden_slides.add(idx)

            texts = []
            has_canvas = False
            for sh_shape in slide.shapes:
                walk_text(sh_shape, texts)
                if sh_shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    w = getattr(sh_shape, "width", 0)
                    h = getattr(sh_shape, "height", 0)
                    if sw and sh and (w >= 0.8 * sw and h >= 0.8 * sh):
                        has_canvas = True

            slide_media = media_map.get(idx, [])
            if has_canvas and not any(t.strip() for t in texts) and slide_media:
                for m in slide_media:
                    canvas_media.add(m)
                    m_path = files.get(m)
                    if m_path:
                        ocr_txt = inspect_pptx.run_ocr_on_image(m_path)
                        if ocr_txt:
                            texts.append(ocr_txt)

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                texts.append(slide.notes_slide.notes_text_frame.text)
            corpus.extend(texts)

            if not include_hidden and is_hidden:
                continue

            slides[idx] = {
                "title": slide.shapes.title.text.strip() if slide.shapes.title else "",
                "media": slide_media,
            }

    visible_media = {m for s in slides.values() for m in s["media"]}
    significant = {
        n for n, (w, h) in sizes.items() 
        if min(w, h) >= MIN_SIDE and (include_hidden or n in visible_media)
    }
    return slides, " ".join(corpus).lower(), significant, sizes, canvas_media, hidden_slides, len(prs.slides)


def spec_strings(slide):
    """All human-readable strings of a spec slide."""
    out = []

    def rec(v):
        if isinstance(v, str):
            out.append(v)
        elif isinstance(v, dict):
            for k, x in v.items():
                if k in ("src", "icon", "icon_src", "kind", "media_ref", "svg"):
                    continue
                rec(x)
        elif isinstance(v, list):
            for x in v:
                rec(x)

    for k in ("title", "subtitle", "badge"):
        if slide.get(k):
            out.append(slide[k])
    rec(slide.get("content", {}))
    return out


def spec_media_refs(slide):
    refs = []

    def rec(v):
        if isinstance(v, dict):
            for k, x in v.items():
                if k in ("src", "media_ref", "icon_src", "image", "file") and isinstance(x, str):
                    refs.append(os.path.basename(x))
                else:
                    rec(x)
        elif isinstance(v, list):
            for x in v:
                rec(x)

    rec(slide.get("content", {}))
    return refs


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    strict = "--strict" in sys.argv
    if not args:
        print(__doc__)
        sys.exit(2)

    spec_path = args[0]
    with open(spec_path, encoding="utf-8") as f:
        spec = yaml.safe_load(f)

    base_arg = args[1] if len(args) > 1 else spec.get("baseline")
    if not base_arg:
        print("No baseline given (argument or `baseline:` key) - nothing to check; "
              "treating as a from-scratch deck.")
        sys.exit(0)
    base_path = base_arg if os.path.isabs(base_arg) or len(args) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(spec_path)), base_arg)
    if not os.path.exists(base_path):
        print(f"ERROR: baseline not found: {base_path}")
        sys.exit(1)

    include_hidden = "--include-hidden" in sys.argv
    b_slides, b_text, significant, sizes, canvas_media, hidden_slides, total_slides = baseline_model(
        base_path, include_hidden=include_hidden
    )
    errors, warns, infos = [], [], []
    slides = spec.get("slides", [])
    cuts = spec.get("baseline_cuts", []) or []

    # --- disallow embedding full-slide canvas screenshots -------------
    for s in slides:
        for ref in spec_media_refs(s):
            if ref in canvas_media:
                errors.append(f"spec slide {s.get('slide_number')} embeds full-slide canvas '{ref}' "
                              "as a content image. Embedding whole-slide screenshots in sub-frames is "
                              "prohibited. Reconstruct with native semantic components or redraw "
                              "diagrams as SVG.")

    # --- slide coverage ------------------------------------------------
    carried = {}
    for s in slides:
        refs = s.get("baseline_slides") or []
        if not refs and not s.get("new"):
            errors.append(f"spec slide {s.get('slide_number')} '{s.get('title')}' has no "
                          "baseline_slides and is not marked `new: true`")
        for r in refs:
            carried.setdefault(r, []).append(s.get("slide_number"))
    cut_slides = set()
    for c in cuts:
        if not c.get("reason"):
            errors.append(f"baseline_cuts entry {c} has no reason")
        if "slide" in c:
            cut_slides.add(c["slide"])
    for n, info in b_slides.items():
        if n not in carried and n not in cut_slides:
            errors.append(f"baseline slide {n} '{info['title'] or 'Untitled'}' is not carried by "
                          "any spec slide and is not in baseline_cuts")
    for n in carried:
        if n not in b_slides and n not in hidden_slides:
            errors.append(f"spec references baseline slide {n}, which does not exist "
                          f"(baseline has {total_slides} total slides)")

    # --- media coverage ------------------------------------------------
    used = set()
    redrawn = set()
    for s in slides:
        used.update(spec_media_refs(s))
        redrawn.update(s.get("redrawn_media") or [])
    cut_media = {c["media"] for c in cuts if "media" in c}
    # media of wholly-cut slides need no separate decision
    cut_slide_media = {m for n in cut_slides for m in b_slides.get(n, {}).get("media", [])}
    for name in sorted(significant):
        if name in used or name in redrawn or name in cut_media or name in cut_slide_media:
            continue
        owners = [n for n, i in b_slides.items() if name in i["media"]]
        # If media is a full-slide canvas and its slide is carried, it is considered reconstructed
        if name in canvas_media and any(o in carried for o in owners):
            continue
        w, h = sizes[name]
        errors.append(f"baseline media {name} ({w}x{h}, slide(s) {owners}) is neither embedded, "
                      "listed in a slide's redrawn_media, nor in baseline_cuts")

    # --- diagram slides carried without any visual --------------------
    for s in slides:
        for r in s.get("baseline_slides") or []:
            info = b_slides.get(r)
            if not info or not info["media"]:
                continue
            content = s.get("content", {})
            has_visual = bool(content.get("image") or content.get("diagram") or content.get("svg")
                              or content.get("mermaid") or s.get("redrawn_media")
                              or any(spec_media_refs(s)))
            if not has_visual:
                warns.append(f"spec slide {s.get('slide_number')} carries baseline slide {r} "
                             f"(has {len(info['media'])} media) but shows no image/diagram")

    # --- invented terms ------------------------------------------------
    for s in slides:
        declared = " ".join(s.get("additions") or []).lower()
        novel = set()
        for text in spec_strings(s):
            for m in re.finditer(r"(?<![.!?]\s)(?<!^)\b([A-Z][A-Za-z0-9]{3,})\b", text):
                w = m.group(1)
                if w in STOP:
                    continue
                if w.lower() in b_text or w.lower() in declared:
                    continue
                novel.add(w)
        if novel:
            warns.append(f"slide {s.get('slide_number')}: terms not in baseline and not in "
                         f"`additions`: {', '.join(sorted(novel))}")
        for a in s.get("additions") or []:
            infos.append(f"slide {s.get('slide_number')} adds (not in baseline): {a}")

    # --- SVG arrow color uniformity check -----------------------------
    def check_svg_arrows(svg_content, source_label):
        markers = {}
        for m in re.finditer(r'<marker\s+[^>]*id=[\"\']([^\"\']+)[\"\'][^>]*>.*?<(?:path|polygon)\s+[^>]*fill=[\"\']([^\"\']+)[\"\']', svg_content, re.DOTALL):
            markers[m.group(1)] = m.group(2).strip().upper()
        errs = []
        for el in re.findall(r'<(?:path|line)\s+[^>]*>', svg_content):
            m_marker = re.search(r'marker-(?:end|start)=[\"\']url\(#([^)]+)\)[\"\']', el)
            m_stroke = re.search(r'stroke=[\"\']([^\"\']+)[\"\']', el)
            if m_marker and m_stroke:
                m_id = m_marker.group(1)
                stroke_col = m_stroke.group(1).strip().upper()
                marker_col = markers.get(m_id)
                if marker_col and stroke_col != marker_col and stroke_col != "NONE":
                    errs.append(f"{source_label}: arrow body stroke ({stroke_col}) != arrowhead marker #{m_id} fill ({marker_col}). "
                                "Arrows must have the same color across their parts.")
        return errs

    # --- SVG arrow collision & obstacle avoidance check --------------
    def check_svg_collisions(svg_content, source_label):
        errs = []
        xml_text = re.sub(r' xmlns="[^"]+"', '', svg_content)
        try:
            root = ET.fromstring(xml_text)
        except Exception:
            return []

        rects = []
        paths = []

        def walk(elem, cur_tx, cur_ty):
            tx, ty = cur_tx, cur_ty
            trans = elem.get("transform", "")
            m = re.search(r"translate\(\s*([-\d.]+)\s*,\s*([-\d.]+)\s*\)", trans)
            if m:
                tx += float(m.group(1))
                ty += float(m.group(2))

            if elem.tag == "rect":
                try:
                    x = float(elem.get("x", 0)) + tx
                    y = float(elem.get("y", 0)) + ty
                    w = float(elem.get("width", 0))
                    h = float(elem.get("height", 0))
                    if 70 <= w <= 450 and 30 <= h <= 300:
                        rects.append((x, y, x + w, y + h, elem))
                except ValueError:
                    pass
            elif elem.tag == "path":
                marker = elem.get("marker-end") or elem.get("marker-start")
                d = elem.get("d", "")
                if marker and d:
                    paths.append((d, tx, ty, elem))

            for child in elem:
                walk(child, tx, ty)

        walk(root, 0, 0)

        for d, tx, ty, elem in paths:
            tokens = re.findall(r"([A-Za-z])|([-\d.]+)", d)
            coords = []
            cur_cmd = None
            current_nums = []
            for cmd, num in tokens:
                if cmd:
                    cur_cmd = cmd.upper()
                elif num and cur_cmd in ("M", "L"):
                    current_nums.append(float(num))
                    if len(current_nums) == 2:
                        coords.append((current_nums[0] + tx, current_nums[1] + ty))
                        current_nums = []
            if len(coords) < 2:
                continue

            for i in range(len(coords) - 1):
                x1, y1 = coords[i]
                x2, y2 = coords[i+1]
                seg_len = ((x2 - x1)**2 + (y2 - y1)**2)**0.5
                if seg_len < 10.0:
                    continue

                for bx1, by1, bx2, by2, r_elem in rects:
                    ibx1, iby1, ibx2, iby2 = bx1 + 8, by1 + 8, bx2 - 8, by2 - 8
                    if ibx1 >= ibx2 or iby1 >= iby2:
                        continue
                    inside_count = 0
                    for step in range(1, 10):
                        t = step / 10.0
                        px = x1 + t * (x2 - x1)
                        py = y1 + t * (y2 - y1)
                        if ibx1 <= px <= ibx2 and iby1 <= py <= iby2:
                            inside_count += 1
                    if inside_count >= 2:
                        errs.append(f"{source_label}: arrow segment ({x1:.0f},{y1:.0f})->({x2:.0f},{y2:.0f}) cuts across element box [{bx1:.0f},{by1:.0f}..{bx2:.0f},{by2:.0f}]. "
                                    "Arrows must route through open corridors and never overlap diagram elements.")
                        break
        return errs

    spec_dir = os.path.dirname(os.path.abspath(spec_path))
    checked_svgs = set()
    for s in slides:
        content = s.get("content", {})
        if content.get("svg"):
            errors.extend(check_svg_arrows(content["svg"], f"slide {s.get('slide_number')} inline svg"))
            errors.extend(check_svg_collisions(content["svg"], f"slide {s.get('slide_number')} inline svg"))
        diag = content.get("diagram", {})
        if isinstance(diag, dict):
            if diag.get("svg"):
                errors.extend(check_svg_arrows(diag["svg"], f"slide {s.get('slide_number')} diagram.svg"))
                errors.extend(check_svg_collisions(diag["svg"], f"slide {s.get('slide_number')} diagram.svg"))
            for k in ("src", "file"):
                if diag.get(k) and isinstance(diag[k], str) and diag[k].endswith(".svg"):
                    p = diag[k] if os.path.isabs(diag[k]) else os.path.join(spec_dir, diag[k])
                    if p not in checked_svgs and os.path.exists(p):
                        checked_svgs.add(p)
                        with open(p, "r", encoding="utf-8", errors="ignore") as f:
                            svg_data = f.read()
                            errors.extend(check_svg_arrows(svg_data, os.path.basename(p)))
                            errors.extend(check_svg_collisions(svg_data, os.path.basename(p)))
        img = content.get("image", {})
        if isinstance(img, dict) and img.get("src") and isinstance(img["src"], str) and img["src"].endswith(".svg"):
            p = img["src"] if os.path.isabs(img["src"]) else os.path.join(spec_dir, img["src"])
            if p not in checked_svgs and os.path.exists(p):
                checked_svgs.add(p)
                with open(p, "r", encoding="utf-8", errors="ignore") as f:
                    svg_data = f.read()
                    errors.extend(check_svg_arrows(svg_data, os.path.basename(p)))
                    errors.extend(check_svg_collisions(svg_data, os.path.basename(p)))

    # --- visual weight & peer highlight check -------------------------
    for s in slides:
        cards = s.get("content", {}).get("cards", [])
        if len(cards) >= 3 and s.get("role") != "comparison":
            highlighted = [c for c in cards if c.get("highlight")]
            if len(highlighted) == 1:
                title_badge = f"{highlighted[0].get('title', '')} {highlighted[0].get('badge', '')}".lower()
                standout = {"recommended", "critical", "warning", "winner", "alert", "priority", "selected"}
                if not any(t in title_badge for t in standout):
                    warns.append(f"slide {s.get('slide_number')}: card '{highlighted[0].get('title')}' has `highlight: true` among {len(cards)} peer cards. "
                                 "Ensure highlight is reserved strictly for standout metrics, critical warnings, or winning options "
                                 "per Theme Accents & Visual Weight Hierarchy mandate.")

    hidden_info = f", {len(hidden_slides)} hidden slides ignored" if hidden_slides and not include_hidden else ""
    print(f"Baseline: {os.path.basename(base_path)} ({len(b_slides)} visible slides{hidden_info}, "
          f"{len(significant)} significant images) -> spec: {len(slides)} slides")
    for label, items in (("ERROR", errors), ("WARN", warns), ("INFO", infos)):
        for i in items:
            print(f"  {label}: {i}")
    print(f"\n{len(errors)} error(s), {len(warns)} warning(s).")
    sys.exit(1 if errors or (strict and warns) else 0)


if __name__ == "__main__":
    main()
